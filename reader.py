"""
reader.py
---------
Reads an EXISTING article (.docx from your article agent, .html from WordPress,
.md or .txt) into one common structure so the Rank Math-style checks work the same
for every format.

What each format can and cannot tell us:
  .docx  headings, paragraphs, tables, images (+ alt text if set), hyperlinks if present.
         (Your article agent's Word files drop hyperlinks, so link tests are "n/a" for them.)
  .html  everything above plus image alt text and links.
  .md    headings, paragraphs, images, links.
  .txt   paragraphs only (no headings, images or links).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional


@dataclass
class Article:
    path: str
    fmt: str
    title: str = ""
    headings: list = field(default_factory=list)      # [(level, text)] for subheadings only (levels 2-6)
    extra_h1: list = field(default_factory=list)      # H1s found in the body (WordPress already uses the title as H1)
    paragraphs: list = field(default_factory=list)    # normal paragraphs (used for the long-paragraph test)
    body_text: str = ""                               # everything except the title, in reading order
    images: list = field(default_factory=list)        # [{"alt": str|None}]
    links: list = field(default_factory=list)         # [(href, text)]
    tables: int = 0
    has_toc: bool = False
    links_detectable: bool = True
    images_detectable: bool = True

    @property
    def words(self) -> list:
        return re.findall(r"[^\s]+", self.body_text)

    @property
    def word_count(self) -> int:
        return len(self.words)


TOC_RE = re.compile(r"^\s*(table of contents|contents|in this (guide|article|post))\s*:?\s*$", re.I)


def _clean(t: str) -> str:
    return re.sub(r"\s+", " ", (t or "")).strip()


# ------------------------------------------------------------------ docx
def read_docx(path: Path) -> Article:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    doc = Document(str(path))
    a = Article(path=str(path), fmt="docx")
    parts, seen_title = [], False

    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "tbl":
            tbl = Table(child, doc)
            a.tables += 1
            for row in tbl.rows:
                cells = [_clean(c.text) for c in row.cells]
                parts.append(" ".join(c for c in cells if c))
            continue
        if tag != "p":
            continue
        p = Paragraph(child, doc)
        text = _clean(p.text)
        style = (p.style.name if p.style is not None else "") or ""

        # images inside this paragraph
        for dp in child.xpath(".//wp:docPr"):
            a.images.append({"alt": _clean(dp.get("descr")) or None})
        # hyperlinks inside this paragraph
        for h in child.xpath(".//w:hyperlink"):
            rid = h.get(qn("r:id"))
            href = doc.part.rels[rid].target_ref if rid and rid in doc.part.rels else None
            if href:
                a.links.append((href, _clean("".join(t.text or "" for t in h.xpath(".//w:t")))))

        if not text:
            continue
        if style == "Title" or (style == "Heading 1" and not seen_title and not a.title):
            if not a.title:
                a.title, seen_title = text, True
                continue
        m = re.match(r"Heading (\d)", style)
        if m:
            lvl = int(m.group(1))
            if lvl == 1:
                a.extra_h1.append(text)
            else:
                a.headings.append((lvl, text))
            parts.append(text)
            if TOC_RE.match(text):
                a.has_toc = True
            continue
        parts.append(text)
        if TOC_RE.match(text):
            a.has_toc = True
        if not style.startswith("List"):
            a.paragraphs.append(text)

    a.body_text = "\n".join(parts)
    return a


# ------------------------------------------------------------------ html
def read_html(path: Path) -> Article:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for t in soup(["script", "style", "noscript", "nav", "footer", "aside", "form"]):
        t.decompose()
    root = soup.find("article") or soup.find("main") or soup.body or soup
    a = Article(path=str(path), fmt="html")

    h1s = root.find_all("h1")
    if h1s:
        a.title = _clean(h1s[0].get_text(" "))
        a.extra_h1 = [_clean(h.get_text(" ")) for h in h1s[1:]]
        h1s[0].decompose()
    elif soup.title:
        a.title = _clean(soup.title.get_text())
    for h in root.find_all(re.compile(r"^h[2-6]$")):
        txt = _clean(h.get_text(" "))
        if txt:
            a.headings.append((int(h.name[1]), txt))
    for h in root.find_all("h1"):
        pass
    for p in root.find_all("p"):
        if p.find_parent(["li", "table", "blockquote"]):
            continue
        txt = _clean(p.get_text(" "))
        if txt:
            a.paragraphs.append(txt)
    for img in root.find_all("img"):
        a.images.append({"alt": _clean(img.get("alt")) or None})
    a.images += [{"alt": None, "media": "video"} for _ in root.find_all(["video", "iframe"])]
    for l in root.find_all("a", href=True):
        a.links.append((l["href"], _clean(l.get_text(" "))))
    a.tables = len(root.find_all("table"))
    a.has_toc = bool(root.select(".rank-math-toc, #rank-math-toc, .wp-block-rank-math-toc-block, "
                                 "[class*='table-of-contents'], [id*='table-of-contents'], .ez-toc-container")) \
        or any(TOC_RE.match(_clean(h.get_text(" "))) for h in root.find_all(re.compile(r"^h[1-6]$|^p$")))
    a.body_text = _clean(root.get_text(" "))
    return a


# ------------------------------------------------------------------ markdown / text
def read_markdown(path: Path) -> Article:
    a = Article(path=str(path), fmt="md")
    text = path.read_text(encoding="utf-8", errors="ignore")
    blocks, buf = [], []
    for line in text.splitlines():
        if line.strip():
            buf.append(line)
        elif buf:
            blocks.append(buf)
            buf = []
    if buf:
        blocks.append(buf)

    parts = []
    for blk in blocks:
        first = blk[0].strip()
        m = re.match(r"^(#{1,6})\s+(.*)$", first)
        if m and len(blk) == 1:
            lvl, txt = len(m.group(1)), _clean(m.group(2))
            if lvl == 1 and not a.title:
                a.title = txt
                continue
            if lvl == 1:
                a.extra_h1.append(txt)
            else:
                a.headings.append((lvl, txt))
            parts.append(txt)
            if TOC_RE.match(txt):
                a.has_toc = True
            continue
        joined = _clean(" ".join(blk))
        for alt, _src in re.findall(r"!\[([^\]]*)\]\(([^)]*)\)", joined):
            a.images.append({"alt": _clean(alt) or None})
        for t, href in re.findall(r"(?<!!)\[([^\]]+)\]\(([^)]+)\)", joined):
            a.links.append((href, t))
        plain = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", joined)
        plain = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", plain)
        plain = re.sub(r"[*_`>#|]+", " ", plain)
        plain = _clean(plain)
        if not plain:
            continue
        if first.startswith("|"):
            a.tables += 1
        parts.append(plain)
        if TOC_RE.match(plain):
            a.has_toc = True
        if not re.match(r"^([-*+]|\d+\.)\s", first) and not first.startswith("|"):
            a.paragraphs.append(plain)
    a.body_text = "\n".join(parts)
    return a


def read_text(path: Path) -> Article:
    a = Article(path=str(path), fmt="txt", links_detectable=False, images_detectable=False)
    raw = path.read_text(encoding="utf-8", errors="ignore")
    paras = [_clean(p) for p in re.split(r"\n\s*\n", raw) if _clean(p)]
    if paras and len(paras[0].split()) <= 20:
        a.title, paras = paras[0], paras[1:]
    a.paragraphs = paras
    a.body_text = "\n".join(paras)
    a.has_toc = any(TOC_RE.match(p) for p in paras)
    return a


def read_article(path) -> Article:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".docx":
        a = read_docx(path)
        a.links_detectable = bool(a.links)      # Word files from the article agent carry no hyperlinks
        a.images_detectable = True
        return a
    if ext in (".html", ".htm"):
        return read_html(path)
    if ext in (".md", ".markdown"):
        return read_markdown(path)
    if ext == ".txt":
        return read_text(path)
    raise ValueError(f"Unsupported file type: {path.name} (use .docx, .html, .md or .txt)")
