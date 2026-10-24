"""
Offline tests: no network, no API key.   Run:  python tests/test_offline.py
The model is faked (it obeys the task lists in the prompts, or misbehaves on purpose), so these tests prove the
document editing, validation, density loop and reporting, not the quality of real Groq output.
"""
import base64
import csv
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import checks  # noqa: E402
import doc as docmod  # noqa: E402
import edits  # noqa: E402
import llm  # noqa: E402
import main  # noqa: E402
import reader  # noqa: E402
import strategy  # noqa: E402

FOCUS = "door draft stopper"
SECONDARY = ["foam door seal", "draft blocker for renters"]
SEMANTIC = ["cold air", "uneven floor"]
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")

FILLER = [
    "Cold air sneaks in through the gap at the bottom of a door and makes the whole room feel colder than it should.",
    "Foam seals compress under the door and fill small gaps without any tools, which suits renters who cannot drill.",
    "Brush sweeps work better on uneven floors because the bristles follow the surface of the threshold.",
    "Adhesive strips hold well on smooth painted wood but can lift on rough or dusty surfaces, so cleaning first matters.",
    "A heavier fabric snake sits on the floor and can be moved aside, which is handy for doors that open often.",
    "Measure the gap with a ruler before you buy anything, because a seal that is too thin leaves a draft behind.",
    "Check whether the door swings over a rug, since thick carpet changes which style of seal will actually fit.",
    "Many readers ask how long a seal lasts, and the honest answer depends on how often the door is used.",
]


def paragraphs(n_words, stuffed=False):
    out, words, i = [], 0, 0
    while words < n_words:
        sents = [FILLER[(i + j) % len(FILLER)] for j in range(4)]
        if i == 3:
            sents[2] = "The FrostGuard foam seal is 36 inches long and fits gaps up to 2 inches."
        if stuffed:
            sents[1] = f"A {FOCUS} blocks air, and a {FOCUS} is easy to fit, so any {FOCUS} will do."
        p = " ".join(sents)
        out.append(p)
        words += len(p.split())
        i += 1
    return out


def add_hyperlink(paragraph, url, text):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r_id = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def make_docx(path, n_words=700, stuffed=False, typo=False):
    from docx import Document
    d = Document()
    d.add_heading("Under Door Draft Stoppers: What Actually Works", level=0)
    ps = paragraphs(n_words, stuffed)
    if typo:
        ps[4] = ps[4].replace("the floor", "teh floor")
    d.add_paragraph(ps[0])
    d.add_heading("Why drafts happen", level=1)                    # extra H1 inside the body
    d.add_paragraph(ps[1])
    d.add_paragraph(ps[2])
    d.add_heading("Foam seals for small gaps", level=2)
    d.add_paragraph(ps[3])
    d.add_heading("Brush sweeps for uneven floors", level=4)       # skipped level
    d.add_paragraph(ps[4])
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Product", "Gap size"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "FrostGuard", "up to 2 inches"
    d.add_paragraph("Foam seal 36 inches fits gaps up to 2 inches.", style="List Bullet")
    for p in ps[5:]:
        d.add_paragraph(p)
    lp = d.add_paragraph("Read our ")
    add_hyperlink(lp, "https://example.com/measure", "measure the door gap")
    pic = Path(tempfile.mkdtemp()) / "x.png"
    pic.write_bytes(PNG)
    d.add_picture(str(pic))
    d.save(str(path))


# ------------------------------------------------------------------ fake model
def parse_blocks(prompt):
    out = []
    for m in re.finditer(r"^\[(\d+)\] (TITLE|H\d|P|LI|TABLE)( \(protected\))?: (.*)$", prompt, re.M):
        out.append({"idx": int(m.group(1)), "kind": m.group(2), "protected": bool(m.group(3)), "text": m.group(4)})
    return out


GOOD_TITLE = "Door Draft Stopper: Proven Picks for 2026"
GOOD_META = (f"Looking for a {FOCUS} that fits your door? See how foam seals and brush sweeps handle small gaps and uneven floors, and which to choose.")


def strategy_json(prompt, title=GOOD_TITLE, h1=None, bad_headings=False):
    blocks = parse_blocks(prompt)
    heads, news = [], []
    for b in blocks:
        if b["kind"] == "H2" and "Foam" in b["text"]:
            heads.append({"idx": b["idx"], "level": 2, "text": "Foam Seals: A Door Draft Stopper for Small Gaps"})
        if b["kind"] == "H4":
            heads.append({"idx": b["idx"], "level": 4, "text": "Brush Sweeps for Uneven Floors"})
        if b["kind"] == "H1" or (b["kind"] == "H1"):
            heads.append({"idx": b["idx"], "level": 1, "text": "Why Cold Drafts Happen at the Bottom of a Door"})
    ps = [b for b in blocks if b["kind"] == "P" and not b["protected"]]
    after_cands = ps[6]
    news.append({"before_idx": after_cands["idx"], "level": 3, "text": " ".join(after_cands["text"].split()[:5])})
    if bad_headings:
        heads += [{"idx": 9999, "level": 2, "text": "Heading that does not exist at all"},
                  {"idx": blocks[0]["idx"], "level": 2, "text": "Title block is not a subheading"},
                  {"idx": next(b["idx"] for b in blocks if b["kind"] == "H2"), "level": 2, "text": "Costs only 19 dollars today"}]
        news.append({"before_idx": ps[8]["idx"], "level": 3, "text": "Warranty and returns policy details"})
    return {
        "search_intent": {"primary": "commercial", "what_searcher_wants": ["which style fits my door", "how to measure the gap", "how to fit it"],
                          "intent_match": "partial", "notes": "The article compares styles but never says how to install."},
        "h1": h1 or f"Door Draft Stopper Guide: What Actually Works",
        "seo_title": title, "meta_description": GOOD_META,
        "headings": heads, "new_headings": news,
        "coverage": [{"topic": "How to measure the gap", "status": "covered", "where_idx": next(b["idx"] for b in blocks if b["kind"] == "H2"),
                      "suggested_heading": "How to measure the gap", "questions": ["What size gap does my door have?"]},
                     {"topic": "Installation steps", "status": "missing", "where_idx": None, "suggested_heading": "How to install a door draft stopper",
                      "questions": ["Do I need tools?", "How long does it take?"]}],
        "secondary_keywords": ["door sweep", "foam seal"], "semantic_keywords": ["cold air", "thick carpet"],
    }


class FakeLLM:
    """Obeys the task lists in the edit prompts unless told to misbehave."""

    def __init__(self, mode="obey", title=GOOD_TITLE, bad_headings=False, h1=None):
        self.mode, self.title, self.bad_headings, self.h1 = mode, title, bad_headings, h1
        self.calls = []

    def __call__(self, prompt, model, max_tokens, temperature=0.3):
        if "Fix these problems with the SEO fields" in prompt:
            self.calls.append("repair")
            return json.dumps({"h1": "Door Draft Stopper Guide: What Actually Works", "seo_title": GOOD_TITLE, "meta_description": GOOD_META})
        if "Fix spelling, grammar and punctuation ONLY" in prompt:
            self.calls.append("grammar")
            out = []
            for m in re.finditer(r"^\[(\d+)\] (.+)$", prompt, re.M):
                if "teh " in m.group(2):
                    out.append({"idx": int(m.group(1)), "text": m.group(2).replace("teh ", "the ")})
            return json.dumps({"edits": out})
        if "SEO copy editor" in prompt:
            self.calls.append("edit")
            if self.mode == "ignore":
                return json.dumps({"edits": []})
            edits_out = []
            for m in re.finditer(r"PARAGRAPH \[(\d+)\].*?\nTASKS: (.*?)\nTEXT: (.*?)(?=\n\nPARAGRAPH|\n\nReturn ONE JSON)", prompt, re.S):
                idx, tasks, text = int(m.group(1)), m.group(2), m.group(3).strip()
                for t in re.finditer(r'(include the exact phrase|use the exact phrase one time LESS[^"]*) "([^"]+)"', tasks):
                    kw = t.group(2)
                    if t.group(1).startswith("include"):
                        text = text + f" This is where a {kw} really matters."
                        if self.mode == "numbers":
                            text += " It costs 19 dollars."
                        if self.mode == "names":
                            text += " Acme Corp makes it."
                    else:
                        text = re.sub(re.escape(kw), "this one", text, count=1, flags=re.I)
                edits_out.append({"idx": idx, "text": text})
            return json.dumps({"edits": edits_out})
        self.calls.append("strategy")
        return json.dumps(strategy_json(prompt, self.title, self.h1, self.bad_headings))


def with_llm(fake, fn):
    saved = llm.chat
    llm.chat = fake
    try:
        return fn()
    finally:
        llm.chat = saved


def cfg_(**seo):
    c = main.load_config(ROOT / "config.yaml")
    c["seo"].update(seo)
    return c


def make_repo(fmt="docx", focus=FOCUS, secondary=SECONDARY, semantic=SEMANTIC, **docx_kw):
    root = Path(tempfile.mkdtemp())
    (root / "articles").mkdir()
    shutil.copy(ROOT / "config.yaml", root / "config.yaml")
    name = f"01_Draft_Stoppers.{fmt}"
    if fmt == "docx":
        make_docx(root / "articles" / name, **docx_kw)
    elif fmt == "html":
        ps = paragraphs(700)
        body = "".join(f"<p>{p}</p>" for p in ps[:3]) + "<h2>Foam seals for small gaps</h2>" + "".join(f"<p>{p}</p>" for p in ps[3:5]) + \
               "<h4>Brush sweeps for uneven floors</h4>" + "".join(f"<p>{p}</p>" for p in ps[5:]) + \
               '<p>Read our <a href="https://example.com/x">measure guide</a>.</p><img src="a.jpg" alt="x">'
        (root / "articles" / name).write_text(f"<html><body><article><h1>Under Door Draft Stoppers</h1>{body}</article></body></html>")
    else:
        ps = paragraphs(700)
        md = "# Under Door Draft Stoppers\n\n" + ps[0] + "\n\n## Foam seals for small gaps\n\n" + "\n\n".join(ps[1:4]) + \
             "\n\n#### Brush sweeps for uneven floors\n\n" + "\n\n".join(ps[4:]) + "\n\nRead our [measure guide](https://example.com/x).\n"
        (root / "articles" / name).write_text(md)
    (root / "jobs.csv").write_text("file,focus_keyword,secondary_keywords,semantic_keywords,power_word\n"
                                   f"{name},{focus},{';'.join(secondary)},{';'.join(semantic)},\n")
    return root, name


def run_main(root, fake, cfg=None):
    out = Path(tempfile.mkdtemp())
    main.PAUSE = 0
    saved = main.ROOT
    main.ROOT = out
    try:
        s = with_llm(fake, lambda: main.run(root, cfg=cfg or cfg_(), out_base=out, today="2026-10-24"))
    finally:
        main.ROOT = saved
    return s, out


# ------------------------------------------------------------------ document model
def test_docx_roundtrip_keeps_everything():
    tmp = Path(tempfile.mkdtemp())
    make_docx(tmp / "a.docx")
    d = docmod.load(tmp / "a.docx")
    kinds = [b.kind for b in d.blocks]
    assert kinds[0] == "title" and "table" in kinds and "li" in kinds
    prot = [b for b in d.blocks if b.protected and b.kind != "table"]
    assert len(prot) == 2                                         # the link paragraph and the picture paragraph
    out = docmod.save(d, tmp / "out", "copy")
    a, b = reader.read_article(tmp / "a.docx"), reader.read_article(out)
    assert a.body_text == b.body_text and a.headings == b.headings and len(a.images) == len(b.images) == 1
    assert a.tables == b.tables == 1 and len(b.links) == 1


def test_docx_edits_applied_in_place():
    tmp = Path(tempfile.mkdtemp())
    make_docx(tmp / "a.docx")
    d = docmod.load(tmp / "a.docx")
    h = next(b for b in d.blocks if b.kind == "h" and b.level == 2)
    h.new_text = "A better heading"
    p = next(b for b in d.blocks if b.kind == "p")
    p.new_text = "Replaced paragraph text."
    p.insert_before.append((2, "Inserted section"))
    d.title_block().new_text = "New Title Here"
    out = docmod.save(d, tmp / "out", "copy")
    a = reader.read_article(out)
    assert a.title == "New Title Here" and (2, "A better heading") in a.headings and (2, "Inserted section") in a.headings
    assert "Replaced paragraph text." in a.body_text
    assert len(a.images) == 1 and a.tables == 1                    # untouched content survives
    assert reader.read_article(tmp / "a.docx").title.startswith("Under Door")      # the original file is unchanged


def test_html_and_md_documents():
    for fmt in ("html", "md"):
        root, name = make_repo(fmt)
        d = docmod.load(root / "articles" / name)
        assert d.title_block().text == "Under Door Draft Stoppers"
        h = next(b for b in d.blocks if b.kind == "h" and b.level == 4)
        h.new_level = 3
        h.new_text = "Brush Sweeps"
        first = next(b for b in d.blocks if b.kind == "p")
        first.new_text = "Changed first paragraph."
        link = next(b for b in d.blocks if b.protected and b.kind == "p")
        link.new_text = "should be ignored"
        out = docmod.save(d, Path(tempfile.mkdtemp()), "copy")
        a = reader.read_article(out)
        assert (3, "Brush Sweeps") in a.headings and "Changed first paragraph." in a.body_text
        assert "should be ignored" not in a.body_text and len(a.links) == 1


def test_normalize_levels_fixes_hierarchy():
    root, name = make_repo()
    d = docmod.load(root / "articles" / name)
    notes = strategy.normalize_levels(d)
    levels = [lvl for lvl, _ in d.outline()]
    assert levels[0] == 1 and levels.count(1) == 1                  # one H1 only
    assert all(b - a <= 1 for a, b in zip(levels, levels[1:]))       # no skipped levels
    assert any("Extra H1" in n for n in notes) and any("no skipped levels" in n for n in notes)


# ------------------------------------------------------------------ strategy guard rails
def test_validate_strategy_blocks_bad_suggestions():
    root, name = make_repo()
    d = docmod.load(root / "articles" / name)
    kw = {"focus": FOCUS, "secondary": SECONDARY, "semantic": SEMANTIC, "power_word": None}
    prompt = strategy.build_prompt(d, kw, cfg_())
    plan = strategy.validate_strategy(strategy_json(prompt, bad_headings=True), d, kw, cfg_())
    idxs = [h["idx"] for h in plan["headings"]]
    assert 9999 not in idxs and d.title_block().idx not in idxs                 # unknown idx and the title block are refused
    assert all("19 dollars" not in h["text"] for h in plan["headings"])          # invented number refused
    assert all(h["level"] >= 2 for h in plan["headings"])
    assert all(n["text"] != "Warranty and returns policy details" for n in plan["new_headings"])   # not described by the text below
    assert len(plan["new_headings"]) == 1
    assert plan["intent"]["primary"] == "commercial" and plan["slug"] if "slug" in plan else True
    assert [c["status"] for c in plan["coverage"]] == ["covered", "missing"]


def test_exact_focus_keyword_in_at_most_two_headings():
    root, name = make_repo()
    d = docmod.load(root / "articles" / name)
    kw = {"focus": FOCUS, "secondary": [], "semantic": [], "power_word": None}
    hs = [b for b in d.blocks if b.kind == "h"]
    raw = {"headings": [{"idx": b.idx, "level": 2, "text": f"Every {FOCUS} variant number {i}"} for i, b in enumerate(hs)]}
    plan = strategy.validate_strategy(raw, d, kw, cfg_())
    assert sum(1 for h in plan["headings"] if checks.has_phrase(h["text"], FOCUS)) <= 2


def test_seo_fields_repair_loop_and_fallback():
    root, name = make_repo()
    d = docmod.load(root / "articles" / name)
    kw = {"focus": FOCUS, "secondary": [], "semantic": [], "power_word": None}
    long_title = "Door Draft Stopper Guide That Covers Every Style You Could Possibly Want to Buy"
    fake = FakeLLM(title=long_title)
    plan = with_llm(fake, lambda: strategy.generate_strategy(d, kw, cfg_()))
    assert "repair" in fake.calls and plan["seo_title"] == GOOD_TITLE and not plan["unresolved"]
    assert plan["slug"] == "door-draft-stopper"

    class Stubborn(FakeLLM):
        def __call__(self, prompt, *a, **k):
            if "Fix these problems with the SEO fields" in prompt:
                return json.dumps({"h1": "Something without the keyword", "seo_title": "x" * 80, "meta_description": "short"})
            return super().__call__(prompt, *a, **k)
    st = Stubborn(title=long_title, h1="No keyword in this heading")
    plan = with_llm(st, lambda: strategy.generate_strategy(d, kw, cfg_()))
    assert plan["seo_title"] == FOCUS.title() and checks.has_phrase(plan["h1"], FOCUS)           # safe fallbacks, flagged
    assert any("Title Case" in n for n in plan["notes"]) and any("H1 is built" in n for n in plan["notes"])


def test_power_word_chosen_by_the_user_is_required():
    probs = strategy.seo_problems({"h1": "Door Draft Stopper Guide", "seo_title": GOOD_TITLE, "meta_description": GOOD_META}, 
                                  {"focus": FOCUS, "power_word": "Ultimate", "secondary": [], "semantic": []}, cfg_(), "article text")
    assert any('power word "Ultimate"' in p for p in probs)


# ------------------------------------------------------------------ edit guard rails
def _kw():
    return {"focus": FOCUS, "secondary": SECONDARY, "semantic": SEMANTIC}


def test_validate_edit_guards():
    orig = "Foam seals compress under the door and fill small gaps up to 2 inches without tools, which suits renters."
    task = [{"type": "add", "kw": FOCUS, "idx": 1}]
    good = orig + f" That is why a {FOCUS} works."
    assert edits.validate_edit(orig, good, task, _kw(), cfg_())[0]
    assert edits.validate_edit(orig, good.replace("2 inches", "3 inches"), task, _kw(), cfg_()) == (False, "changed a number")
    assert not edits.validate_edit(orig, good + " Acme sells it.", task, _kw(), cfg_())[0]                    # new proper name
    assert edits.validate_edit(orig, "Totally different sentence about cats and dogs playing outside today now.", task, _kw(), cfg_())[0] is False
    assert edits.validate_edit(orig, orig + " more words" * 12 + f" {FOCUS}.", task, _kw(), cfg_())[1] in ("changed the length too much", "changed the meaning too much")
    assert edits.validate_edit(orig, orig + " and nothing else changes here.", task, _kw(), cfg_()) == (False, f"did not add '{FOCUS}'")
    three = orig + f" A {FOCUS}, a {FOCUS} and a {FOCUS}."
    assert not edits.validate_edit(orig, three, task, _kw(), cfg_())[0]                                      # more than twice
    stuffed = f"A {FOCUS} blocks air and a {FOCUS} fits easily."
    removed = "A door draft stopper blocks air and this one fits easily."
    assert edits.validate_edit(stuffed, removed, [{"type": "remove", "kw": FOCUS, "idx": 1}], _kw(), cfg_())[0]


def test_protected_paragraphs_are_never_candidates():
    root, name = make_repo()
    d = docmod.load(root / "articles" / name)
    cands = edits.candidates(d)
    assert cands and all(not b.protected and b.kind == "p" for b in cands)


# ------------------------------------------------------------------ end to end
def _numbers(text):
    return sorted(strategy._num_norm(t) for t in strategy.NUM.findall(text))


def test_end_to_end_docx_reaches_density_target():
    root, name = make_repo(typo=True)
    original = (root / "articles" / name).read_bytes()
    fake = FakeLLM()
    s, out = run_main(root, fake)
    r = s["results"][0]
    assert s["status"] == "success" and r["density_before"] < 1.0
    assert abs(r["density_after"] - 1.7) <= 0.2, r
    base = out / "01-draft-stoppers"
    for f in ("optimized_article.docx", "rank_math_fields.md", "heading_structure.md", "topical_coverage.md", "changes.md", "seo_report.md", "seo_package.json"):
        assert (base / f).exists(), f
    assert (root / "articles" / name).read_bytes() == original                       # your file is untouched

    before, after = reader.read_article(root / "articles" / name), reader.read_article(base / "optimized_article.docx")
    assert after.title.startswith(FOCUS.title()) or checks.has_phrase(after.title, FOCUS)
    levels = [1] + [lvl for lvl, _ in after.headings]
    assert after.extra_h1 == [] and all(b - a <= 1 for a, b in zip(levels, levels[1:]))        # clean H1 > H2 > H3 hierarchy
    assert sum(1 for _, t in after.headings if checks.has_phrase(t, FOCUS)) <= 2
    assert checks.has_phrase(checks.first_10_percent(after.words), FOCUS)
    assert all(checks.count_phrase(after.body_text, k) >= 1 for k in SECONDARY + SEMANTIC)
    assert _numbers(before.body_text) == _numbers(after.body_text)                    # no number added or lost
    assert len(after.images) == 1 and after.tables == 1 and len(after.links) == 1     # image, table, link kept
    assert after.word_count > before.word_count

    fields = (base / "rank_math_fields.md").read_text()
    assert f"{FOCUS}, {SECONDARY[0]}, {SECONDARY[1]}" in fields and GOOD_TITLE in fields and "door-draft-stopper" in fields
    pkg = json.loads((base / "seo_package.json").read_text())
    t = {x["test"]: x["status"] for x in pkg["checks_after"]}
    assert t["Focus keyword in SEO title"] == t["Focus keyword in URL"] == t["Focus keyword in subheadings (H2-H6)"] == "pass"
    assert t["Keyword density (focus + secondary keywords)"] == "pass"
    assert not any("image" in x["test"].lower() or "link" in x["test"].lower() for x in pkg["checks_after"])      # skipped on request
    changes = (base / "changes.md").read_text()
    assert "**Before**" in changes and "**After**" in changes
    assert (out / "summary.csv").exists()


def test_stuffed_article_is_brought_down():
    root, name = make_repo(stuffed=True)
    s, out = run_main(root, FakeLLM())
    r = s["results"][0]
    assert r["density_before"] > 2.5 and r["density_after"] < r["density_before"]
    assert abs(r["density_after"] - 1.7) <= 0.4


def test_refused_rewrites_never_reach_the_file():
    for mode in ("numbers", "names"):
        root, name = make_repo()
        s, out = run_main(root, FakeLLM(mode=mode))
        base = out / "01-draft-stoppers"
        text = reader.read_article(base / "optimized_article.docx").body_text
        assert "19 dollars" not in text and "Acme" not in text
        assert json.loads((base / "seo_package.json").read_text())["rejected_edits"]
        assert "refused" in (base / "changes.md").read_text()
        assert abs(s["results"][0]["density_after"] - 1.7) > 0.2           # and the report says honestly it did not get there
        assert "Not on target" in (base / "seo_report.md").read_text()


def test_model_that_ignores_tasks_is_reported_honestly():
    root, name = make_repo()
    s, out = run_main(root, FakeLLM(mode="ignore"))
    assert s["status"] == "success"
    rep = (out / "01-draft-stoppers" / "seo_report.md").read_text()
    assert "Density reached" in rep
    assert (out / "01-draft-stoppers" / "optimized_article.docx").exists()      # headings and SEO fields are still delivered


def test_html_and_md_end_to_end():
    for fmt, ext in (("html", "html"), ("md", "md")):
        root, name = make_repo(fmt)
        s, out = run_main(root, FakeLLM())
        assert s["status"] == "success", s
        a = reader.read_article(out / "01-draft-stoppers" / f"optimized_article.{ext}")
        levels = [1] + [lvl for lvl, _ in a.headings]
        assert all(b - a_ <= 1 for a_, b in zip(levels, levels[1:])) and len(a.links) == 1
        assert abs(s["results"][0]["density_after"] - 1.7) <= 0.2
        assert checks.has_phrase(checks.first_10_percent(a.words), FOCUS)


def test_suggested_keywords_when_none_given():
    root, name = make_repo(secondary=[], semantic=[])
    s, out = run_main(root, FakeLLM())
    assert s["status"] == "success"
    rep = (out / "01-draft-stoppers" / "seo_report.md").read_text()
    assert "suggested: confirm" in rep and "door sweep" in rep


def test_optional_grammar_pass():
    root, name = make_repo(typo=True)
    s, out = run_main(root, FakeLLM(), cfg_(grammar=True))
    assert "teh " not in reader.read_article(out / "01-draft-stoppers" / "optimized_article.docx").body_text
    root2, name2 = make_repo(typo=True)
    s2, out2 = run_main(root2, FakeLLM(), cfg_(grammar=False))
    assert "teh " in reader.read_article(out2 / "01-draft-stoppers" / "optimized_article.docx").body_text      # off by default


def test_no_api_key_gives_audit_only():
    root, name = make_repo()

    def no_key(*a, **k):
        raise llm.LLMUnavailable("GROQ_API_KEY is not set")
    s, out = run_main(root, no_key)
    assert s["results"][0]["status"].startswith("audit only")
    assert "No optimized copy was made" in (out / "01-draft-stoppers" / "seo_report.md").read_text()
    assert not (out / "01-draft-stoppers" / "optimized_article.docx").exists()


def test_missing_focus_keyword_and_missing_file_fail_cleanly():
    root, name = make_repo(focus="")
    s, out = run_main(root, FakeLLM())
    assert "focus_keyword" in s["results"][0]["status"]
    root2, _ = make_repo()
    (root2 / "jobs.csv").write_text(f"file,focus_keyword\nnope.docx,{FOCUS}\n")
    s2, _ = run_main(root2, FakeLLM())
    assert "not found" in s2["results"][0]["status"]


def test_parse_list_and_jobs():
    assert main.parse_list("Foam Door Seal; draft blocker for renters |foam door seal") == ["foam door seal", "draft blocker for renters"]
    root, name = make_repo()
    j = main.load_jobs(root)[0]
    assert j["focus_keyword"] == FOCUS and j["secondary"] == SECONDARY and j["semantic"] == SEMANTIC


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"PASS  {name}")
        except Exception as e:  # noqa: BLE001
            failed += 1
            import traceback
            print(f"FAIL  {name}: {e!r}")
            traceback.print_exc()
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    sys.exit(1 if failed else 0)
