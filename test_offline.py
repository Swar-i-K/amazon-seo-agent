"""
Offline tests: no network, no API key.   Run:  python tests/test_offline.py
The fake LLM below returns canned JSON, so these tests prove the reading, checking,
validation, repair and export logic, not the quality of real Groq answers.
"""
import base64
import csv
import json
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import checks  # noqa: E402
import exporters  # noqa: E402
import keywords as kw  # noqa: E402
import llm  # noqa: E402
import main  # noqa: E402
import reader  # noqa: E402
import seo  # noqa: E402

KW = "best under door draft stopper that stays in place"
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")

FILLER = [
    "Cold air sneaks in through the gap at the bottom of a door and makes the whole room feel colder than it should.",
    "Foam seals compress under the door and fill small gaps without any tools, which suits renters who cannot drill.",
    "Brush sweeps work better on uneven floors because the bristles follow the surface of the threshold.",
    "Adhesive strips hold well on smooth painted wood but can lift on rough or dusty surfaces, so cleaning first matters.",
    "A heavier fabric snake sits on the floor and can be moved aside, which is handy for doors that open often.",
    "Measure the gap with a ruler before you buy anything, because a seal that is too thin leaves a draft behind.",
    "Check whether the door swings over a rug, since thick carpet changes which style of seal will actually fit.",
    "Many readers ask how long a seal lasts, and the honest answer depends on how often the door is used.",
]


def body_paragraphs(n_words, kw_places=()):
    paras, words, i = [], 0, 0
    while words < n_words:
        sents = [FILLER[(i + j) % len(FILLER)] for j in range(4)]
        if i in kw_places:
            sents[1] = f"This guide explains the {KW} so you can pick one fast."
        p = " ".join(sents)
        paras.append(p)
        words += len(p.split())
        i += 1
    return paras


def add_hyperlink(paragraph, url, text):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r_id = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                    is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def make_docx(path, kw_in_intro=True, with_links=False, n_words=700, subheading_kw=False):
    from docx import Document
    d = Document()
    d.add_heading("Under Door Draft Stoppers: What Actually Works", level=0)
    paras = body_paragraphs(n_words, kw_places=(0,) if kw_in_intro else ())
    d.add_paragraph(paras[0])
    d.add_heading(f"{KW.title()} for rough doors" if subheading_kw else "Foam seals for small gaps", level=2)
    for p in paras[1:3]:
        d.add_paragraph(p)
    d.add_heading("Brush sweeps for uneven floors", level=2)
    d.add_paragraph("Foam seal 36 inches fits gaps up to 2 inches.", style="List Bullet")
    t = d.add_table(rows=2, cols=3)
    for i, v in enumerate(["Product", "Gap size", "Install"]):
        t.rows[0].cells[i].text = v
    for i, v in enumerate(["FrostGuard foam", "up to 2 inches", "Adhesive strip"]):
        t.rows[1].cells[i].text = v
    for p in paras[3:]:
        d.add_paragraph(p)
    pic = Path(tempfile.mkdtemp()) / "x.png"
    pic.write_bytes(PNG)
    d.add_picture(str(pic))
    if with_links:
        lp = d.add_paragraph("Read our ")
        add_hyperlink(lp, "https://example.com/blog/measure-door-gap", "measure the door gap")
        add_hyperlink(lp, "https://other.org/seal", "external guide")
    d.save(str(path))


HTML = """<html><head><title>x</title></head><body><article>
<h1>Draft Stopper Guide</h1>
<div class="rank-math-toc">Contents</div>
<p>%s</p>
<h2>Best under door draft stopper that stays in place</h2>
<p>%s</p>
<h3>Foam</h3>
<img src="a.jpg" alt="foam draft stopper on a door"><img src="b.jpg">
<p><a href="https://example.com/blog/measure">measure your gap</a> and see <a href="https://other.org/x">this test</a>
and <a href="/blog/relative">relative link</a></p>
</article></body></html>""" % (" ".join(body_paragraphs(120, (0,))), " ".join(body_paragraphs(60)))


# ------------------------------------------------------------------ reader
def test_read_docx():
    tmp = Path(tempfile.mkdtemp())
    make_docx(tmp / "a.docx", with_links=True)
    a = reader.read_article(tmp / "a.docx")
    assert a.title == "Under Door Draft Stoppers: What Actually Works"
    assert [h for h in a.headings] == [(2, "Foam seals for small gaps"), (2, "Brush sweeps for uneven floors")]
    assert a.tables == 1 and len(a.images) == 1 and a.word_count > 600
    assert len(a.links) == 2 and a.links_detectable
    assert all(len(p.split()) < 120 for p in a.paragraphs)
    assert "Foam seal 36 inches" not in " ".join(a.paragraphs)       # list items are not paragraphs


def test_docx_without_links_is_not_testable():
    tmp = Path(tempfile.mkdtemp())
    make_docx(tmp / "a.docx")
    a = reader.read_article(tmp / "a.docx")
    rows = checks.run_checks(a, KW, [], None, None, None, main.load_config())
    link_rows = [r for r in rows if r["test"] in ("Internal link", "External link")]
    assert link_rows and all(r["status"] == "n/a" for r in link_rows)


def test_read_html():
    tmp = Path(tempfile.mkdtemp())
    (tmp / "a.html").write_text(HTML)
    a = reader.read_article(tmp / "a.html")
    assert a.title == "Draft Stopper Guide" and a.has_toc
    assert [h[0] for h in a.headings] == [2, 3]
    assert [i["alt"] for i in a.images] == ["foam draft stopper on a door", None]
    assert len(a.links) == 3
    cfg = main.load_config()
    cfg["site"]["domain"] = "https://example.com"
    rows = {r["test"]: r for r in checks.run_checks(a, KW, [], None, None, None, cfg)}
    assert rows["Internal link"]["status"] == "pass" and "2 internal" in rows["Internal link"]["detail"]
    assert rows["External link"]["status"] == "pass"
    assert rows["Focus keyword in image alt text"]["status"] == "fail"       # the alt text lacks the keyword


def test_read_markdown_and_txt():
    tmp = Path(tempfile.mkdtemp())
    (tmp / "a.md").write_text("# Title here\n\nIntro with the keyword door draft.\n\n## Sub one\n\n"
                              "Text with [a link](https://x.com/y) and ![alt text](i.png)\n\n| a | b |\n|---|---|\n| 1 | 2 |\n")
    a = reader.read_article(tmp / "a.md")
    assert a.title == "Title here" and a.headings == [(2, "Sub one")] and a.tables == 1
    assert a.links == [("https://x.com/y", "a link")] and a.images == [{"alt": "alt text"}]
    (tmp / "a.txt").write_text("Short title\n\n" + "word " * 60 + "\n\n" + "more " * 30)
    t = reader.read_article(tmp / "a.txt")
    assert t.title == "Short title" and len(t.paragraphs) == 2 and not t.images_detectable
    try:
        reader.read_article(tmp / "a.pdf")
        assert False
    except ValueError:
        pass


# ------------------------------------------------------------------ checks
def _art(**kw_):
    tmp = Path(tempfile.mkdtemp())
    make_docx(tmp / "a.docx", **kw_)
    return reader.read_article(tmp / "a.docx")


def test_phrase_counting():
    assert checks.count_phrase("Best under-door draft stopper. The best under door draft stopper!", "best under door draft stopper") == 2
    assert checks.count_phrase("bestunder door", "best under door") == 0
    assert checks.slugify("Best Under Door Draft Stopper That Stays in Place") == "best-under-door-draft-stopper-that-stays-in-place"


def test_first_10_percent():
    inside = _art(kw_in_intro=True)
    outside = _art(kw_in_intro=False)
    cfg = main.load_config()
    r1 = {r["test"]: r for r in checks.run_checks(inside, KW, [], None, None, None, cfg)}
    r2 = {r["test"]: r for r in checks.run_checks(outside, KW, [], None, None, None, cfg)}
    assert r1["Focus keyword in first 10% of content"]["status"] == "pass"
    assert r2["Focus keyword in first 10% of content"]["status"] == "fail"
    assert r2["Focus keyword in content"]["status"] == "fail"


def test_content_length_and_subheading_checks():
    cfg = main.load_config()
    short = {r["test"]: r for r in checks.run_checks(_art(n_words=300), KW, [], None, None, None, cfg)}
    assert short["Content length (600+ words)"]["status"] == "fail"
    longer = {r["test"]: r for r in checks.run_checks(_art(n_words=700, subheading_kw=True), KW, [], None, None, None, cfg)}
    assert longer["Content length (600+ words)"]["status"] == "pass"
    assert longer["Focus keyword in subheadings (H2-H6)"]["status"] == "pass"
    no = {r["test"]: r for r in checks.run_checks(_art(), KW, [], None, None, None, cfg)}
    assert no["Focus keyword in subheadings (H2-H6)"]["status"] == "fail"


def test_density_bands():
    a = _art(kw_in_intro=True)
    cfg = main.load_config()
    base = {r["test"]: r for r in checks.run_checks(a, KW, [], None, None, None, cfg)}["Keyword density (primary + secondary)"]
    assert base["status"] == "warn"                                           # one mention in ~700 words: low
    a.body_text += (" " + KW) * 40                                           # stuffing
    stuffed = {r["test"]: r for r in checks.run_checks(a, KW, [], None, None, None, cfg)}["Keyword density (primary + secondary)"]
    assert stuffed["status"] == "fail"
    a2 = _art()
    n = round(len(a2.words) * 0.012)                                         # ~1.2%
    a2.body_text += (" " + KW + " filler words here to pad out the count") * n
    ok = {r["test"]: r for r in checks.run_checks(a2, KW, [], None, None, None, cfg)}["Keyword density (primary + secondary)"]
    assert ok["status"] in ("pass", "warn")


def test_title_and_meta_checks():
    a, cfg = _art(), main.load_config()
    good_t = "Best Under Door Draft Stopper That Stays in Place (2026)"
    good_m = (f"Looking for the {KW}? See how foam, brush and adhesive seals compare and which one fits your door.")
    assert len(good_t) <= 60 and 110 <= len(good_m) <= 160
    rows = {r["test"]: r for r in checks.run_checks(a, KW, [], good_t, good_m, checks.slugify(KW), cfg)}
    for t in ("Focus keyword in SEO title", "Focus keyword in meta description", "Focus keyword in URL",
              "Focus keyword at the beginning of the SEO title", "Number in SEO title", "Power word in SEO title"):
        assert rows[t]["status"] == "pass", t
    bad = {r["test"]: r for r in checks.run_checks(a, KW, [], "How to Pick a Draft Stopper for Your Front Door This Winter and Save Money", "short", "x", cfg)}
    assert bad["Focus keyword in SEO title"]["status"] == "fail" and bad["Focus keyword in URL"]["status"] == "fail"
    assert bad["Focus keyword at the beginning of the SEO title"]["status"] == "fail"
    assert any(r["test"].startswith("SEO title length") and r["status"] == "fail" for r in bad.values())
    probs = checks.metadata_problems(KW, "x" * 70, "short", cfg)
    assert any("70 characters" in p for p in probs) and any("START" in p for p in probs)


def test_long_paragraph_and_toc():
    a, cfg = _art(), main.load_config()
    a.paragraphs.append("word " * 130)
    rows = {r["test"]: r for r in checks.run_checks(a, KW, [], None, None, None, cfg)}
    assert rows["Short paragraphs (120 words or fewer)"]["status"] == "fail"
    assert rows["Table of contents"]["status"] == "fail"
    a.has_toc = True
    assert {r["test"]: r for r in checks.run_checks(a, KW, [], None, None, None, cfg)}["Table of contents"]["status"] == "pass"


# ------------------------------------------------------------------ validation guard rails
def test_validate_core_blocks_bad_output():
    a, cfg = _art(), main.load_config()
    pages = [{"url": "https://example.com/blog/measure-door-gap", "title": "Measure a door gap"}]
    raw = {
        "seo_title": "T", "meta_description": "M",
        "secondary_keywords": ["door draft stopper for renters", "pizza delivery near me", KW],
        "heading_fixes": [
            {"current": "Foam seals for small gaps", "suggested": f"{KW.title()}: foam seals", "reason": "kw"},
            {"current": "A heading that does not exist", "suggested": "whatever", "reason": "x"},
            {"current": "Brush sweeps for uneven floors", "suggested": f"The {KW} on uneven floors", "reason": "kw"},
            {"current": "Foam seals for small gaps", "suggested": f"Another {KW} take", "reason": "third kw"},
        ],
        "intro_fix": {"current_sentence": "This sentence is not in the article at all.", "suggested": f"The {KW}."},
        "faqs": [
            {"question": "How wide a gap does foam cover?", "answer": "Foam seals fit gaps up to 2 inches."},
            {"question": "How much does it cost?", "answer": "It costs $14.99 on average."},              # number not in article
            {"question": "No question mark", "answer": "x"},
        ],
        "comparison_table": {"worth_adding": True, "reason": "r", "caption": "c", "columns": ["Product", "Gap", "Price"],
                             "rows": [["FrostGuard", "up to 2 inches", "$18.99"], ["SnugSeal", "up to 3 inches", "Not stated"]]},
        "internal_links": [
            {"url": "https://example.com/blog/measure-door-gap", "anchor_text": "measure the gap", "section": "s"},
            {"url": "https://evil.example/x", "anchor_text": "ruler before you buy", "section": "s"},          # not allowed
            {"url": "https://example.com/blog/measure-door-gap", "anchor_text": "Foam seals compress", "section": "s"},  # dup
        ],
        "image_alts": [{"image_number": 1, "alt": f"{KW} installed"}, {"image_number": 5, "alt": "no such image"}],
    }
    v = seo.validate_core(raw, a, pages, KW, cfg)
    assert v["secondary_keywords"] == ["door draft stopper for renters"]
    assert [f["current"] for f in v["heading_fixes"]] == ["Foam seals for small gaps", "Brush sweeps for uneven floors"]
    assert v["intro_fix"] is None
    assert [f["question"] for f in v["faqs"]] == ["How wide a gap does foam cover?"]
    ct = v["comparison_table"]
    assert ct["rows"][0][1] == "up to 2 inches" and ct["rows"][0][2] == "Not verified"     # invented number replaced
    assert ct["rows"][1][1] == "Not verified"                                              # 3 inches not in article
    assert [l["anchor_text"] for l in v["internal_links"]] == ["measure the gap"] or v["internal_links"] == []
    assert v["image_alts"] == [{"image_number": 1, "alt": f"{KW} installed"}]
    assert any("not found in the article" in n for n in v["notes"])


def test_validate_extras_limits():
    a = _art()
    raw = {"pins": [
        {"angle": "a", "title": "T" * 120, "description": "d", "image_direction": "i"},                 # title too long
        {"angle": "b", "title": "Stop the draft", "description": "Cold air under the door? Read on.", "image_direction": "door"},
        {"angle": "b", "title": "Same angle again", "description": "d", "image_direction": "i"},       # duplicate angle
        {"angle": "c", "title": "Under $5 deal", "description": "Only $5!", "image_direction": "i"},   # invented number
    ], "facebook": [
        {"style": "humor", "text": "My door has a draft and an attitude."},
        {"style": "humor", "text": "Second joke should be dropped."},
        {"style": "tip", "text": "Measure the gap first. [article link]"},
        {"style": "tip", "text": "See https://x.com now"},                                              # link not allowed
    ]}
    v = seo.validate_extras(raw, a, KW)
    assert [p["angle"] for p in v["pins"]] == ["b"]
    assert [f["style"] for f in v["facebook"]] == ["humor", "tip"]


# ------------------------------------------------------------------ repair loop + fake LLM
GOOD_TITLE = "Best Under Door Draft Stopper That Stays in Place (2026)"
GOOD_META = f"Looking for the {KW}? See how foam, brush and adhesive seals compare and which one fits your door."


def core_json(title=GOOD_TITLE, meta=GOOD_META):
    return {
        "seo_title": title, "meta_description": meta,
        "secondary_keywords": ["door draft stopper for renters", "under door seal"],
        "keyword_map": [{"keyword": "door draft stopper for renters", "intent": "commercial", "place": "H2", "note": "n"}],
        "heading_fixes": [{"current": "Foam seals for small gaps", "suggested": f"{KW.title()}: foam seals", "reason": "kw"}],
        "intro_fix": {"current_sentence": FILLER[0], "suggested": f"The {KW} stops cold air."},
        "faqs": [{"question": "How wide a gap does foam cover?", "answer": "Foam seals fit gaps up to 2 inches."}],
        "comparison_table": {"worth_adding": True, "reason": "Two products compared.", "caption": "Seals",
                             "columns": ["Product", "Gap size"], "rows": [["FrostGuard foam", "up to 2 inches"], ["Brush sweep", "Not stated"]]},
        "internal_links": [{"url": "https://example.com/blog/measure-door-gap", "anchor_text": "measure the gap", "section": "Foam"}],
        "image_alts": [{"image_number": 1, "alt": "foam seal under a white door"}],
    }


class FakeLLM:
    def __init__(self, first_title=None, fail_repair=False, focus="best under door draft stopper that stays in place"):
        self.calls, self.first_title, self.fail_repair, self.focus = [], first_title, fail_repair, focus

    def __call__(self, prompt, model, max_tokens, temperature=0.3):
        if "Pick the single best Rank Math focus keyword" in prompt:
            self.calls.append("focus")
            return json.dumps({"focus_keyword": self.focus})
        if "Fix these problems" in prompt:
            self.calls.append("repair")
            return json.dumps({"seo_title": "still too long " * 6 if self.fail_repair else GOOD_TITLE, "meta_description": GOOD_META})
        if "social media editor" in prompt:
            self.calls.append("extras")
            return json.dumps({"pins": [
                {"angle": f"angle {i}", "title": f"Pin title {i}", "description": "Cold air under the door? Read the guide.",
                 "image_direction": "A white door with a foam seal"} for i in range(3)],
                "facebook": [{"style": "humor", "text": "My door has a draft and an attitude."},
                             {"style": "tip", "text": "Measure the gap first. [article link]"},
                             {"style": "checklist", "text": "Gap, floor type, door traffic. [article link]"}]})
        self.calls.append("core")
        return json.dumps(core_json(self.first_title or GOOD_TITLE))


def _with_llm(fake, fn):
    saved = llm.chat
    llm.chat = fake
    try:
        return fn()
    finally:
        llm.chat = saved


def test_repair_loop_fixes_long_title():
    a, cfg = _art(), main.load_config()
    fake = FakeLLM(first_title="Best Under Door Draft Stopper That Stays in Place and Beats Every Other Option")
    core = _with_llm(fake, lambda: seo.generate_core(a, KW, [], [], cfg))
    assert "repair" in fake.calls and core["seo_title"] == GOOD_TITLE and not core["unresolved"]
    assert core["slug"] == checks.slugify(KW)


def test_repair_failure_uses_fallback_and_flags():
    a, cfg = _art(), main.load_config()
    fake = FakeLLM(first_title="How to choose a draft stopper for your door in winter time to save heat", fail_repair=True)
    core = _with_llm(fake, lambda: seo.generate_core(a, KW, [], [], cfg))
    assert core["seo_title"] == KW.title()
    assert any("Title Case" in n for n in core["notes"])


# ------------------------------------------------------------------ end to end
def _make_repo(with_jobs=True, focus=KW, with_pages=True):
    root = Path(tempfile.mkdtemp())
    (root / "articles").mkdir()
    shutil.copy(ROOT / "config.yaml", root / "config.yaml")
    make_docx(root / "articles" / "01_Draft_Stoppers.docx", kw_in_intro=False)
    (root / "jobs.csv").write_text("file,focus_keyword\n" + (f"01_Draft_Stoppers.docx,{focus}\n" if with_jobs else ""))
    (root / "site_pages.csv").write_text("url,title\n" + ("https://example.com/blog/measure-door-gap,Measure a door gap\n" if with_pages else ""))
    return root


def test_end_to_end():
    root, out = _make_repo(), Path(tempfile.mkdtemp())
    main.PAUSE = 0
    s = _with_llm(FakeLLM(), lambda: main.run(root, out_base=out, today="2026-10-23"))
    assert s["status"] == "success" and s["processed"] == 1
    base = out / "01-draft-stoppers"
    for f in ("seo_report.md", "rank_math_fields.md", "faq.html", "comparison_table.html", "internal_links.csv",
              "social.md", "seo_package.json"):
        assert (base / f).exists(), f
    fields = (base / "rank_math_fields.md").read_text()
    assert GOOD_TITLE in fields and checks.slugify(KW) in fields and f"{KW}, door draft stopper for renters" in fields
    rep = (base / "seo_report.md").read_text()
    assert "Your article was not changed" in rep and "Intro fix" in rep           # keyword missing from intro -> fix shown
    assert "Not verified" in rep and "After fixes" in rep
    pkg = json.loads((base / "seo_package.json").read_text())
    assert pkg["check_summary"]["total"] > 10
    t = {r["test"]: r["status"] for r in pkg["checks"]}
    assert t["Focus keyword in SEO title"] == "pass" and t["Focus keyword in URL"] == "pass"
    after = {r["test"]: r["status"] for r in pkg["checks_after_fixes"]}
    assert t["Focus keyword in subheadings (H2-H6)"] == "fail" and after["Focus keyword in subheadings (H2-H6)"] == "pass"
    assert t["Focus keyword in first 10% of content"] == "fail" and after["Focus keyword in first 10% of content"] == "pass"
    assert after["Focus keyword in image alt text"] == "fail"          # fake alt text has no keyword: only one alt may carry it
    assert pkg["seo"] and "_primary" not in pkg["seo"]
    assert (out / "summary.csv").exists()
    assert "<table>" in (base / "comparison_table.html").read_text()
    # the source article is untouched
    assert (root / "articles" / "01_Draft_Stoppers.docx").exists()


def test_focus_keyword_suggested_when_blank():
    root, out = _make_repo(focus=""), Path(tempfile.mkdtemp())
    main.PAUSE = 0
    orig, kw.autocomplete = kw.autocomplete, (lambda q, timeout=10: [])
    try:
        fake = FakeLLM()
        s = _with_llm(fake, lambda: main.run(root, out_base=out, today="2026-10-23"))
    finally:
        kw.autocomplete = orig
    assert s["status"] == "success" and "focus" in fake.calls
    assert "please confirm" in (out / "01-draft-stoppers" / "seo_report.md").read_text()


def test_no_api_key_still_gives_content_audit():
    root, out = _make_repo(), Path(tempfile.mkdtemp())
    main.PAUSE = 0

    def no_key(*a, **k):
        raise llm.LLMUnavailable("GROQ_API_KEY is not set")
    s = _with_llm(no_key, lambda: main.run(root, out_base=out, today="2026-10-23"))
    assert s["status"] == "success"
    rep = (out / "01-draft-stoppers" / "seo_report.md").read_text()
    assert "LLM step skipped" in rep and "Rank Math checklist" in rep
    assert not (out / "01-draft-stoppers" / "faq.html").exists()


def test_blank_keyword_and_no_llm_fails_cleanly():
    root, out = _make_repo(focus=""), Path(tempfile.mkdtemp())
    main.PAUSE = 0
    orig, kw.autocomplete = kw.autocomplete, (lambda q, timeout=10: [])

    def no_key(*a, **k):
        raise llm.LLMUnavailable("GROQ_API_KEY is not set")
    try:
        s = _with_llm(no_key, lambda: main.run(root, out_base=out, today="2026-10-23"))
    finally:
        kw.autocomplete = orig
    assert s["status"] == "failed" and "focus_keyword" in s["results"][0]["status"]


def test_missing_file_and_unreadable_file():
    root, out = _make_repo(), Path(tempfile.mkdtemp())
    main.PAUSE = 0
    (root / "jobs.csv").write_text(f"file,focus_keyword\nnope.docx,{KW}\nbad.docx,{KW}\n01_Draft_Stoppers.docx,{KW}\n")
    (root / "articles" / "bad.docx").write_text("not really a docx")
    s = _with_llm(FakeLLM(), lambda: main.run(root, out_base=out, today="2026-10-23"))
    assert s["status"] == "partial" and s["processed"] == 1
    assert [r["status"].split(":")[0] for r in s["results"]] == ["failed", "failed", "success"]


def test_apply_suggestions_does_not_touch_original():
    a = _art(kw_in_intro=False)
    original_body, original_heads = a.body_text, list(a.headings)
    core = core_json()
    core["_primary"] = KW
    core["heading_fixes"] = [{"current": "Foam seals for small gaps", "suggested": f"{KW.title()}: foam seals", "reason": "r"}]
    core["intro_fix"] = {"current_sentence": FILLER[0], "suggested": f"The {KW} stops cold air."}
    core["image_alts"] = [{"image_number": 1, "alt": f"{KW} under a door"}]
    b = seo.apply_suggestions(a, core)
    assert a.body_text == original_body and a.headings == original_heads            # original untouched
    assert checks.has_phrase(b.body_text, KW) and any(checks.has_phrase(t, KW) for _l, t in b.headings)
    assert b.images[0]["alt"] == f"{KW} under a door"


def test_jobs_default_to_all_files():
    root = _make_repo(with_jobs=False)
    assert [j["file"] for j in main.load_jobs(root)] == ["01_Draft_Stoppers.docx"]


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"PASS  {name}")
        except Exception as e:  # noqa: BLE001
            failed += 1
            import traceback
            print(f"FAIL  {name}: {e!r}")
            traceback.print_exc()
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    sys.exit(1 if failed else 0)

